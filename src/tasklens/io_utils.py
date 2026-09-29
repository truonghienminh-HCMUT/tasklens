"""Đọc đề bài từ PDF / DOCX / TXT / MD thành văn bản thuần."""
from __future__ import annotations

import io
from pathlib import Path


def extract_text(data: bytes, filename: str) -> str:
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if suffix == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts).strip()
    return data.decode("utf-8", errors="replace").replace("\r\n", "\n").strip()


def read_file(path: str | Path) -> str:
    path = Path(path)
    return extract_text(path.read_bytes(), path.name)
